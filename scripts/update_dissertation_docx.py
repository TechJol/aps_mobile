#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches


SOURCE_DOCX = "/Users/mamarasulov/Downloads/Диссертация (New) (1).docx"


def find_paragraph_index(document: Document, prefix: str) -> int:
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip().startswith(prefix):
            return index
    raise ValueError(f"Paragraph starting with {prefix!r} was not found")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def add_paragraph_before(target, text: str = "", style: str | None = None):
    paragraph = target.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def add_table_before(target, rows: list[list[str]]) -> None:
    document = target._parent
    table = document.add_table(rows=len(rows), cols=len(rows[0]), width=Inches(6.3))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = str(value)
    target._element.addprevious(table._element)


def add_picture_before(target, image_path: Path, width_inches: float = 6.2) -> None:
    paragraph = add_paragraph_before(target)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def format_float(value: object, digits: int = 4) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


def build_model_metrics_table(metrics: pd.DataFrame) -> list[list[str]]:
    rows = [["Модель", "Accuracy", "Precision", "Recall", "Average Precision", "ROC AUC"]]
    names = {
        "logistic_regression": "Logistic Regression",
        "random_forest": "Random Forest",
    }
    for _, row in metrics.iterrows():
        rows.append(
            [
                names.get(row["model"], row["model"]),
                format_float(row["accuracy"]),
                format_float(row["precision"]),
                format_float(row["recall"]),
                format_float(row["average_precision"]),
                format_float(row["roc_auc"]),
            ]
        )
    return rows


def build_scenario_table(summary: pd.DataFrame) -> list[list[str]]:
    rows = [
        [
            "ID",
            "Категория",
            "Подход",
            "Precision",
            "Hit-rate",
            "Ожидаемые тесты",
            "Test cases",
            "Время, c",
            "Статус",
        ]
    ]
    for _, row in summary.iterrows():
        rows.append(
            [
                row["scenario_id"],
                row["category"],
                "ML" if row["approach"] == "ml" else "Baseline",
                format_float(row["precision"]),
                format_float(row["hit_rate"]),
                f"{int(row['expected_hits'])}/{int(row['expected_total'])}",
                str(int(row["test_case_count"])),
                format_float(row["duration_s"], 2),
                row["status"],
            ]
        )
    return rows


def build_aggregate_table(summary: pd.DataFrame) -> list[list[str]]:
    aggregate = (
        summary.groupby("approach")
        .agg(
            avg_precision=("precision", "mean"),
            avg_hit_rate=("hit_rate", "mean"),
            total_expected_hits=("expected_hits", "sum"),
            total_expected=("expected_total", "sum"),
            avg_duration_s=("duration_s", "mean"),
        )
        .reset_index()
    )
    rows = [["Подход", "Средняя precision", "Средняя hit-rate", "Ожидаемые тесты", "Среднее время, c"]]
    for _, row in aggregate.iterrows():
        rows.append(
            [
                "ML" if row["approach"] == "ml" else "Baseline",
                format_float(row["avg_precision"]),
                format_float(row["avg_hit_rate"]),
                f"{int(row['total_expected_hits'])}/{int(row['total_expected'])}",
                format_float(row["avg_duration_s"], 2),
            ]
        )
    return rows


def add_updated_practice_content(document: Document, target) -> None:
    metrics = pd.read_csv("scripts/out/ml_training_metrics.csv")
    summary = pd.read_csv("scripts/out/scenarios/scenario_summary.csv")
    figures_dir = Path("scripts/out/figures")

    add_paragraph_before(target, "2.6. Проектирование модели интеллектуальной приоритизации тестов", "Heading 2")
    add_paragraph_before(
        target,
        "В условиях частых изменений программного кода полный запуск всех тестов при каждом обновлении может быть избыточным. "
        "Для проекта SoftkgPro была спроектирована модель интеллектуальной приоритизации тестов, основанная на анализе связи между измененными файлами и тестовыми сценариями. "
        "Задача модели заключается в построении ранжированного списка тестов, которые целесообразно выполнять в первую очередь после изменения кода.",
    )
    add_paragraph_before(
        target,
        "Входными данными модели являются пути измененных файлов, структура модулей приложения, список существующих тестов и экспертная оценка критичности тестов. "
        "Выходными данными являются ранги тестов и оценка релевантности пары «измененный файл - тест».",
    )
    add_paragraph_before(target, "Рисунок 2.5 - Pipeline ML-приоритизации тестов")
    add_picture_before(target, figures_dir / "fig_01_ml_pipeline.png", 6.3)
    add_paragraph_before(
        target,
        "Как показано на рисунке 2.5, экспериментальный pipeline включает сбор данных, извлечение признаков, обучение модели, ранжирование тестов, запуск выбранных проверок и формирование отчетов.",
    )
    add_paragraph_before(target, "Таблица 2.4 - Признаки, использованные для обучения модели")
    add_table_before(
        target,
        [
            ["Группа признаков", "Примеры", "Назначение"],
            ["Структурные", "same_feature, same_layer", "Оценивают архитектурную близость файла и теста"],
            ["Лексические", "token_jaccard, common_token_count", "Оценивают сходство путей и имен модулей"],
            ["Тип теста", "is_unit_test, is_widget_test", "Учитывают уровень тестирования"],
            ["Критичность", "manual_priority", "Учитывают экспертный приоритет сценария"],
            ["Инфраструктурные", "changed_is_network, changed_is_routes", "Помогают выбирать тесты при изменениях общего кода"],
        ],
    )

    add_paragraph_before(target, "2.7. Реализация механизма отбора и ранжирования тестов", "Heading 2")
    add_paragraph_before(
        target,
        "Практическая реализация выполнена в виде набора Python-скриптов в каталоге scripts. "
        "Скрипт build_ml_dataset.py формирует датасет из пар «измененный файл - тест», train_test_prioritizer.py обучает модели, "
        "predict_tests_ml.py выполняет ML-ранжирование, run_experiment_scenarios.py проверяет сценарии изменений, а run_ml_experiment.py объединяет весь эксперимент в один воспроизводимый запуск.",
    )
    add_paragraph_before(
        target,
        "Для сравнения был сохранен baseline-подход, основанный на эвристическом совпадении токенов в путях файлов. "
        "Это позволило оценить, дает ли ML-модель практическое преимущество по сравнению с простым структурным правилом.",
    )
    add_paragraph_before(
        target,
        "В итоговом эксперименте было проанализировано 297 Dart-файлов и 12 тестовых файлов. "
        "Сформированный датасет содержит 3564 строки, из которых 457 пар размечены как релевантные, а 3107 - как нерелевантные.",
    )

    add_paragraph_before(target, "2.8. Практическая апробация разработанного подхода", "Heading 2")
    add_paragraph_before(
        target,
        "Практическая апробация проводилась на шести сценариях изменений: auth, payment, income, menu, core/network и routes/widgets. "
        "Для каждого сценария были определены измененные файлы и ожидаемые тесты. Затем каждый сценарий обрабатывался двумя подходами: ML и baseline. "
        "Для каждого подхода выбирались Top-5 тестовых файлов и выполнялся фактический запуск командой flutter test.",
    )
    add_paragraph_before(target, "Таблица 2.5 - Результаты обучения моделей")
    add_table_before(target, build_model_metrics_table(metrics))
    add_paragraph_before(target, "Рисунок 2.6 - Сравнение качества обученных моделей")
    add_picture_before(target, figures_dir / "fig_02_model_metrics.png", 6.3)
    add_paragraph_before(
        target,
        "По результатам обучения лучшей моделью стала Random Forest. Она достигла accuracy 0.9618, precision 0.7703, recall 1.0000, average precision 0.9830 и ROC AUC 0.9975.",
    )

    add_paragraph_before(target, "2.9. Сравнительный анализ результатов тестирования", "Heading 2")
    add_paragraph_before(
        target,
        "Сравнительный анализ выполнялся по метрикам precision, hit-rate, количеству найденных ожидаемых тестов, числу выполненных test cases и времени выполнения. "
        "Hit-rate показывает, какую долю ожидаемых тестов удалось включить в Top-5 выбранных тестов.",
    )
    add_paragraph_before(target, "Таблица 2.6 - Сравнение ML и baseline по сценариям")
    add_table_before(target, build_scenario_table(summary))
    add_paragraph_before(target, "Рисунок 2.7 - Hit-rate по сценариям изменений")
    add_picture_before(target, figures_dir / "fig_03_hit_rate_by_scenario.png", 6.3)
    add_paragraph_before(target, "Рисунок 2.8 - Количество найденных ожидаемых тестов")
    add_picture_before(target, figures_dir / "fig_04_expected_hits_total.png", 6.3)
    add_paragraph_before(
        target,
        "В агрегированном сравнении ML-подход нашел 12 из 13 ожидаемых тестов, тогда как baseline нашел 8 из 13. "
        "Наиболее выраженное преимущество ML наблюдается в сценариях изменения инфраструктурных частей приложения: core/network и routes/widgets.",
    )

    add_paragraph_before(target, "2.10. Оценка эффективности предложенного подхода", "Heading 2")
    add_paragraph_before(target, "Таблица 2.7 - Агрегированные результаты сценарного эксперимента")
    add_table_before(target, build_aggregate_table(summary))
    add_paragraph_before(target, "Рисунок 2.9 - Время выполнения выбранных тестов")
    add_picture_before(target, figures_dir / "fig_05_duration_by_scenario.png", 6.3)
    add_paragraph_before(target, "Рисунок 2.10 - Количество выполненных test cases")
    add_picture_before(target, figures_dir / "fig_06_test_cases_by_scenario.png", 6.3)
    add_paragraph_before(
        target,
        "Средняя hit-rate для ML составила 0.9167, для baseline - 0.6667. Средняя precision для ML составила 0.4000, для baseline - 0.2667. "
        "Среднее время выполнения выбранных тестов для ML составило 5.24 секунды, для baseline - 4.46 секунды. "
        "Разница во времени объясняется тем, что ML чаще выбирал более релевантные widget/smoke-проверки, которые не всегда являются самыми короткими.",
    )
    add_paragraph_before(
        target,
        "Полученные результаты показывают, что предложенный ML-подход повышает полноту выбора релевантных тестов по сравнению с baseline-эвристикой. "
        "При этом все сценарные test runs завершились успешно, что подтверждает исполнимость выбранных наборов тестов и практическую применимость решения.",
    )
    add_paragraph_before(
        target,
        "Ограничением эксперимента является небольшой размер тестовой базы и использование структурно-экспертной разметки вместо большой истории CI-дефектов. "
        "В дальнейшем модель может быть усилена данными о прошлых падениях тестов, анализом импортов, call graph и историей уязвимостей.",
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Create an updated dissertation DOCX with ML practice section.")
    parser.add_argument("--source", default=SOURCE_DOCX)
    parser.add_argument("--output", default="scripts/out/Диссертация_updated_ml_practice.docx")
    args = parser.parse_args()

    document = Document(args.source)
    start = find_paragraph_index(document, "2.6.")
    end = find_paragraph_index(document, "2.11.")

    for paragraph in list(document.paragraphs[start:end]):
        remove_paragraph(paragraph)

    target = document.paragraphs[find_paragraph_index(document, "2.11.")]
    add_updated_practice_content(document, target)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Updated DOCX: {output.as_posix()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
